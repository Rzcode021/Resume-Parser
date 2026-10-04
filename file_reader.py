from pypdf import PdfReader
from docx import Document


# ==================================
# READ PDF FILE
# ==================================

def read_pdf(uploaded_file):

    reader = PdfReader(uploaded_file)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:

            text += page_text + "\n"

    return text


# ==================================
# READ DOCX FILE
# ==================================

def read_docx(uploaded_file):

    document = Document(uploaded_file)

    text = ""


    # Read paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text += paragraph.text + "\n"


    # Read tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text += cell.text + "\n"

    return text


# ==================================
# DETECT FILE TYPE
# ==================================

def read_resume(uploaded_file):

    file_name = uploaded_file.name.lower()


    if file_name.endswith(".pdf"):

        return read_pdf(uploaded_file)


    elif file_name.endswith(".docx"):

        return read_docx(uploaded_file)


    else:

        raise ValueError(
            "Unsupported file type. Please upload PDF or DOCX."
        )